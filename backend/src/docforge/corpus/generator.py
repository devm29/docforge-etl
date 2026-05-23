"""Generates the bundled demo corpus: invoices and contracts with realistic imperfections.

Every quirk here is deliberate. Some invoices print US-style dates and currency symbols that a
strict schema rejects (exercising the repair loop), one has an arithmetic error the vendor made
(no repair can fix it, so it lands in review), and one says "Due upon receipt" instead of a date.
"""

from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import docx
from docx.shared import Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.styles import Font

CUSTOMER = "Acme Robotics, Inc."
FIXED_TIMESTAMP = datetime(2026, 4, 15, 9, 0, tzinfo=UTC)
CUSTOMER_ADDRESS = ["1200 Harbor Way", "Oakland, CA 94607"]


@dataclass(frozen=True)
class LineItem:
    description: str
    quantity: Decimal
    unit_price: Decimal

    @property
    def amount(self) -> Decimal:
        return (self.quantity * self.unit_price).quantize(Decimal("0.01"))


@dataclass(frozen=True)
class InvoiceSpec:
    filename: str
    vendor: str
    vendor_address: tuple[str, ...]
    number: str
    issued: date
    due: date | None
    items: tuple[LineItem, ...]
    tax_rate: Decimal
    currency: str = "USD"
    us_format: bool = False
    total_override: Decimal | None = None
    notes: str = ""
    po_number: str | None = None

    @property
    def subtotal(self) -> Decimal:
        return sum((item.amount for item in self.items), Decimal("0.00"))

    @property
    def tax(self) -> Decimal:
        return (self.subtotal * self.tax_rate).quantize(Decimal("0.01"))

    @property
    def total(self) -> Decimal:
        return self.total_override if self.total_override is not None else self.subtotal + self.tax

    def money(self, value: Decimal) -> str:
        if self.us_format:
            symbol = "$" if self.currency == "USD" else f"{self.currency} "
            return f"{symbol}{value:,.2f}"
        return f"{value:.2f}"

    def day(self, value: date) -> str:
        return value.strftime("%B %-d, %Y") if self.us_format else value.isoformat()


def D(value: str) -> Decimal:  # noqa: N802
    return Decimal(value)


INVOICES: tuple[InvoiceSpec, ...] = (
    InvoiceSpec(
        filename="invoices/2026-03/northwind-INV-24017.pdf",
        vendor="Northwind Office Supply",
        vendor_address=("88 Market Street, Suite 400", "San Francisco, CA 94105"),
        number="INV-24017",
        issued=date(2026, 3, 3),
        due=date(2026, 4, 2),
        items=(
            LineItem("Ergonomic task chair, mesh back", D("12"), D("289.00")),
            LineItem("Height-adjustable desk frame", D("6"), D("540.00")),
            LineItem("Recycled copy paper, 10-ream case", D("20"), D("48.50")),
            LineItem("Assembly and delivery", D("1"), D("350.00")),
        ),
        tax_rate=D("0.0875"),
        notes="Payment terms: Net 30. Please reference the invoice number on your remittance.",
        po_number="PO-88213",
    ),
    InvoiceSpec(
        filename="invoices/2026-03/helios-cloud-HC-1180.pdf",
        vendor="Helios Cloud Services",
        vendor_address=("2100 Mission College Blvd", "Santa Clara, CA 95054"),
        number="HC-1180",
        issued=date(2026, 3, 1),
        due=date(2026, 3, 31),
        items=(
            LineItem("Compute reserved instances, March", D("1"), D("18420.00")),
            LineItem("Object storage, 84 TB-month", D("84"), D("21.00")),
            LineItem("Premium support plan", D("1"), D("2500.00")),
        ),
        tax_rate=D("0"),
        us_format=True,
        notes="Usage details are available in the Helios console under Billing > Statements.",
    ),
    InvoiceSpec(
        filename="invoices/2026-03/brightline-BL-55302.pdf",
        vendor="Brightline Logistics",
        vendor_address=("4410 Port Road", "Long Beach, CA 90802"),
        number="BL-55302",
        issued=date(2026, 3, 9),
        due=date(2026, 4, 8),
        items=(
            LineItem("Freight, Oakland to Reno, full truckload", D("3"), D("1875.00")),
            LineItem("Liftgate service", D("3"), D("95.00")),
            LineItem("Fuel surcharge", D("1"), D("412.40")),
        ),
        tax_rate=D("0"),
        us_format=True,
        total_override=D("6497.40"),
        notes="Detention charges, if any, are billed separately.",
    ),
    InvoiceSpec(
        filename="invoices/2026-03/quantum-print-QP-0931.pdf",
        vendor="Quantum Print Co.",
        vendor_address=("19 Foundry Lane", "Emeryville, CA 94608"),
        number="QP-0931",
        issued=date(2026, 3, 12),
        due=None,
        items=(
            LineItem("Trade show banners, 3x8 ft", D("4"), D("165.00")),
            LineItem("Product brochures, 1,000 count", D("2"), D("420.00")),
        ),
        tax_rate=D("0.0925"),
        notes="Due upon receipt. Thank you for your business.",
    ),
    InvoiceSpec(
        filename="invoices/2026-04/atlas-facilities-AF-7710.pdf",
        vendor="Atlas Facilities Management",
        vendor_address=("600 Grand Avenue", "Oakland, CA 94610"),
        number="AF-7710",
        issued=date(2026, 4, 1),
        due=date(2026, 4, 30),
        items=(
            LineItem("Janitorial services, April", D("1"), D("6800.00")),
            LineItem("HVAC preventive maintenance", D("2"), D("975.00")),
            LineItem("Window cleaning, exterior", D("1"), D("1450.00")),
            LineItem("Restroom supplies replenishment", D("1"), D("312.75")),
        ),
        tax_rate=D("0"),
        po_number="PO-88340",
    ),
    InvoiceSpec(
        filename="invoices/2026-04/bluebird-catering-BC-2044.pdf",
        vendor="Bluebird Catering",
        vendor_address=("71 Linden Street", "Oakland, CA 94607"),
        number="BC-2044",
        issued=date(2026, 4, 6),
        due=date(2026, 4, 20),
        items=(
            LineItem("All-hands lunch, 140 guests", D("140"), D("24.50")),
            LineItem("Coffee and pastry service", D("140"), D("7.25")),
            LineItem("Event staff, 4 hours", D("6"), D("180.00")),
        ),
        tax_rate=D("0.1025"),
        us_format=True,
    ),
    InvoiceSpec(
        filename="invoices/2026-04/cedar-pine-legal-CP-3307.pdf",
        vendor="Cedar & Pine Legal LLP",
        vendor_address=("500 Capitol Mall, Floor 12", "Sacramento, CA 95814"),
        number="CP-3307",
        issued=date(2026, 4, 2),
        due=date(2026, 5, 2),
        items=tuple(
            LineItem(description, hours, D("425.00"))
            for description, hours in (
                ("Review of supplier master services agreement", D("3.5")),
                ("Negotiation call with Helios counsel", D("1.5")),
                ("Drafting data processing addendum", D("4.0")),
                ("Trademark clearance search, ROVER mark", D("2.5")),
                ("Board consent preparation", D("1.0")),
                ("Employment agreement updates, 3 hires", D("3.0")),
                ("Export control classification memo", D("5.5")),
                ("Lease amendment review, Harbor Way", D("2.0")),
                ("Correspondence with landlord counsel", D("0.5")),
                ("Privacy policy refresh", D("3.0")),
                ("NDA template revisions", D("1.0")),
                ("Contractor IP assignment review", D("1.5")),
                ("Patent filing coordination", D("2.5")),
                ("Regulatory filing review", D("2.0")),
                ("Vendor dispute strategy session", D("1.5")),
                ("Insurance policy review", D("1.0")),
                ("Stock option plan amendment", D("3.5")),
                ("Securities law compliance check", D("2.0")),
                ("Customer contract redlines, Orion Labs", D("4.5")),
                ("Data breach tabletop exercise", D("3.0")),
                ("Open-source license audit", D("2.5")),
                ("Annual compliance certification", D("1.0")),
                ("Warranty terms update", D("1.5")),
                ("General corporate advice", D("2.0")),
                ("Meeting with CFO regarding audit", D("1.0")),
                ("Contract template library refresh", D("2.5")),
                ("Board meeting attendance", D("3.0")),
                ("Records retention policy", D("1.5")),
            )
        ),
        tax_rate=D("0"),
        us_format=True,
        notes="Hours are billed in six-minute increments. Disbursements are itemized separately.",
    ),
)


@dataclass(frozen=True)
class Clause:
    heading: str
    paragraphs: tuple[str, ...]


@dataclass(frozen=True)
class ContractSpec:
    filename: str
    title: str
    intro: str
    clauses: tuple[Clause, ...]
    signatories: tuple[tuple[str, str, str], ...] = field(default=())


CONTRACTS: tuple[ContractSpec, ...] = (
    ContractSpec(
        filename="contracts/helios-master-services-agreement.docx",
        title="Master Services Agreement",
        intro=(
            'This Master Services Agreement (the "Agreement") is entered into as of January 15, 2026 '
            '(the "Effective Date") by and between Helios Cloud Services, Inc. ("Provider") and '
            'Acme Robotics, Inc. ("Customer").'
        ),
        clauses=(
            Clause(
                "1. Services",
                (
                    "Provider will make the hosted compute, storage and support services described in each "
                    "Order Form available to Customer during the Term.",
                    "Provider will maintain availability of at least 99.9% measured monthly, excluding "
                    "scheduled maintenance announced at least 72 hours in advance.",
                ),
            ),
            Clause(
                "2. Term and Renewal",
                (
                    "The initial term of this Agreement is twenty-four (24) months from the Effective Date and "
                    "ends on January 14, 2028.",
                    "This Agreement will automatically renew for successive twelve-month periods unless either "
                    "party gives sixty (60) days' prior written notice of non-renewal.",
                ),
            ),
            Clause(
                "3. Fees and Payment",
                (
                    "Customer will pay total fees of $486,000.00 over the initial term, invoiced monthly in "
                    "arrears. Amounts are stated in USD.",
                    "Invoices are payable within thirty (30) days. Late amounts accrue interest at 1% per month.",
                ),
            ),
            Clause(
                "4. Confidentiality",
                (
                    "Each party will protect the other party's Confidential Information with at least the "
                    "same degree of care it uses for its own, and no less than reasonable care.",
                ),
            ),
            Clause(
                "5. Limitation of Liability",
                (
                    "Except for breaches of Section 4, neither party's aggregate liability will exceed the fees "
                    "paid in the twelve months preceding the claim.",
                ),
            ),
            Clause(
                "6. Governing Law",
                (
                    "This Agreement is governed by the laws of the State of California, without regard to its "
                    "conflict of laws principles.",
                ),
            ),
        ),
        signatories=(
            ("Helios Cloud Services, Inc.", "Dana Whitfield", "Chief Revenue Officer"),
            ("Acme Robotics, Inc.", "Priya Raman", "Chief Financial Officer"),
        ),
    ),
    ContractSpec(
        filename="contracts/orion-labs-mutual-nda.docx",
        title="Mutual Non-Disclosure Agreement",
        intro=(
            'This Mutual Non-Disclosure Agreement (the "Agreement") is made as of February 2, 2026 by and '
            'between Orion Labs LLC ("Orion") and Acme Robotics, Inc. ("Acme").'
        ),
        clauses=(
            Clause(
                "1. Purpose",
                (
                    "The parties wish to evaluate a potential joint development of warehouse navigation "
                    "software and may disclose Confidential Information for that purpose.",
                ),
            ),
            Clause(
                "2. Obligations",
                (
                    "The receiving party will use Confidential Information only for the purpose above and will "
                    "not disclose it to third parties without prior written consent.",
                ),
            ),
            Clause(
                "3. Term",
                (
                    "This Agreement remains in effect until February 1, 2028. It shall not automatically renew.",
                    "Either party may terminate this Agreement with thirty (30) days' prior written notice.",
                ),
            ),
            Clause(
                "4. Governing Law",
                ("This Agreement is governed by the laws of the State of Delaware.",),
            ),
        ),
        signatories=(
            ("Orion Labs LLC", "Marcus Bell", "Managing Member"),
            ("Acme Robotics, Inc.", "Priya Raman", "Chief Financial Officer"),
        ),
    ),
    ContractSpec(
        filename="contracts/brightline-transportation-agreement.docx",
        title="Transportation Services Agreement",
        intro=(
            'This Transportation Services Agreement (the "Agreement") is entered into as of March 1, 2026 by '
            'and between Brightline Logistics Corp. ("Carrier") and Acme Robotics, Inc. ("Shipper").'
        ),
        clauses=(
            Clause(
                "1. Scope",
                (
                    "Carrier will transport Shipper's goods between the locations listed in Schedule A using "
                    "equipment that meets all applicable safety regulations.",
                ),
            ),
            Clause(
                "2. Rates",
                (
                    "Rates are set out in Schedule A. The estimated contract value for the initial term is "
                    "$120,000.00, which is not a minimum commitment.",
                ),
            ),
            Clause(
                "3. Term",
                (
                    "The initial term of this Agreement is twelve (12) months. It will automatically renew for "
                    "one-year terms unless terminated with ninety (90) days' prior written notice.",
                ),
            ),
            Clause(
                "4. Governing Law",
                ("This Agreement is governed by the laws of the State of Nevada.",),
            ),
        ),
    ),
)

LEASE = ContractSpec(
    filename="contracts/harbor-way-lease-amendment.pdf",
    title="First Amendment to Office Lease",
    intro=(
        'This First Amendment to Office Lease (the "Amendment") is entered into as of April 1, 2026 '
        '(the "Effective Date") by and between Harbor Way Properties LP ("Landlord") and '
        'Acme Robotics, Inc. ("Tenant").'
    ),
    clauses=(
        Clause(
            "1. Extension of Term",
            (
                "The term of the Lease is extended by thirty-six (36) months and will expire on March 31, 2029.",
            ),
        ),
        Clause(
            "2. Base Rent",
            (
                "Base rent during the extended term totals $1,296,000.00, payable in equal monthly installments "
                "of $36,000.00 on the first day of each month.",
            ),
        ),
        Clause(
            "3. Renewal",
            (
                "Tenant may extend the term for one additional period of five years by giving Landlord one "
                "hundred eighty (180) days' prior written notice. The Lease will not automatically renew.",
            ),
        ),
        Clause(
            "4. Governing Law",
            ("This Amendment is governed by the laws of the State of California.",),
        ),
    ),
    signatories=(
        ("Harbor Way Properties LP", "Elena Ortiz", "General Partner"),
        ("Acme Robotics, Inc.", "Priya Raman", "Chief Financial Officer"),
    ),
)


def _invoice_pdf(spec: InvoiceSpec, path: Path) -> None:
    pdf = FPDF(unit="pt", format="letter")
    pdf.set_auto_page_break(auto=True, margin=72)
    pdf.set_margins(54, 54, 54)
    pdf.set_title(f"Invoice {spec.number}")
    pdf.set_author(spec.vendor)
    pdf.set_creation_date(FIXED_TIMESTAMP)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(0, 26, spec.vendor, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    for line in spec.vendor_address:
        pdf.cell(0, 12, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(18)

    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 20, "INVOICE", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 10)
    meta = [
        f"Invoice Number: {spec.number}",
        f"Invoice Date: {spec.day(spec.issued)}",
        f"Due Date: {spec.day(spec.due)}" if spec.due else "Payment Terms: Due upon receipt",
        f"Currency: {spec.currency}",
    ]
    if spec.po_number:
        meta.append(f"PO Number: {spec.po_number}")
    for line in meta:
        pdf.cell(0, 14, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(14)

    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 15, "Bill To", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    for line in (CUSTOMER, *CUSTOMER_ADDRESS):
        pdf.cell(0, 13, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(16)

    pdf.set_font("Helvetica", "", 9)
    with pdf.table(
        col_widths=(260, 60, 90, 94),
        text_align=("LEFT", "RIGHT", "RIGHT", "RIGHT"),
        line_height=16,
        first_row_as_headings=True,
        repeat_headings=1,
    ) as table:
        table.row(("Description", "Qty", "Unit Price", "Amount"))
        for item in spec.items:
            quantity = f"{item.quantity.normalize():f}"
            table.row((item.description, quantity, spec.money(item.unit_price), spec.money(item.amount)))
    pdf.ln(14)

    pdf.set_font("Helvetica", "", 10)
    tax_label = f"Tax ({(spec.tax_rate * 100).normalize():f}%)" if spec.tax_rate else "Tax"
    for label, value in (("Subtotal", spec.subtotal), (tax_label, spec.tax)):
        pdf.cell(0, 15, f"{label}: {spec.money(value)}", align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 18, f"Total Due: {spec.money(spec.total)}", align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    if spec.notes:
        pdf.ln(22)
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(0, 12, spec.notes)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


def _contract_docx(spec: ContractSpec, path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = spec.title
    document.core_properties.author = "Acme Robotics Legal"
    document.core_properties.created = FIXED_TIMESTAMP.replace(tzinfo=None)
    document.core_properties.modified = FIXED_TIMESTAMP.replace(tzinfo=None)
    style = document.styles["Normal"]
    style.font.size = Pt(11)
    document.add_heading(spec.title, level=0)
    document.add_paragraph(spec.intro)
    for clause in spec.clauses:
        document.add_heading(clause.heading, level=1)
        for paragraph in clause.paragraphs:
            document.add_paragraph(paragraph)
    if spec.signatories:
        document.add_heading("Signatures", level=1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ("Party", "Name", "Title"), strict=True):
            cell.text = label
        for party, name, title in spec.signatories:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = party, name, title
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    _freeze_zip(path)


def _contract_pdf(spec: ContractSpec, path: Path) -> None:
    pdf = FPDF(unit="pt", format="letter")
    pdf.set_margins(72, 72, 72)
    pdf.set_title(spec.title)
    pdf.set_creation_date(FIXED_TIMESTAMP)
    pdf.add_page()
    pdf.set_font("Times", "B", 18)
    pdf.multi_cell(0, 24, spec.title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(8)
    pdf.set_font("Times", "", 11)
    pdf.multi_cell(0, 15, spec.intro, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    for clause in spec.clauses:
        pdf.ln(10)
        pdf.set_font("Times", "B", 13)
        pdf.cell(0, 18, clause.heading, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Times", "", 11)
        for paragraph in clause.paragraphs:
            pdf.multi_cell(0, 15, paragraph, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(4)
    if spec.signatories:
        pdf.ln(16)
        pdf.set_font("Times", "B", 13)
        pdf.cell(0, 18, "Signatures", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Times", "", 10)
        with pdf.table(col_widths=(200, 120, 148), line_height=16) as table:
            table.row(("Party", "Name", "Title"))
            for signatory in spec.signatories:
                table.row(signatory)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


def _invoice_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Invoice"
    issued = date(2026, 4, 10)
    rows: list[tuple[object, ...]] = [
        ("Sierra Components Ltd.",),
        (),
        ("Invoice Number:", "SC-90021"),
        ("Invoice Date:", issued),
        ("Due Date:", issued + timedelta(days=45)),
        ("Currency:", "USD"),
        ("Bill To:", CUSTOMER),
        (),
        ("Line Items",),
        ("Description", "Qty", "Unit Price", "Amount"),
        ("Brushless DC motor, 48V", 40, 86.5, 3460.0),
        ("Motor controller board", 40, 54.25, 2170.0),
        ("Harness kit, 1.2 m", 80, 6.4, 512.0),
        (),
        ("Subtotal:", 6142.0),
        ("Tax:", 0.0),
        ("Total Due:", 6142.0),
    ]
    for row in rows:
        sheet.append(list(row))
    sheet["A1"].font = Font(bold=True, size=14)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))
    _freeze_zip(path)


def _freeze_zip(path: Path) -> None:
    """openpyxl stamps the save time into the package, which would change the hash on every build."""
    stamp = FIXED_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, payload in entries:
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                    rb"\g<1>" + stamp.encode() + rb"\g<2>",
                    payload,
                )
            info = zipfile.ZipInfo(name, date_time=FIXED_TIMESTAMP.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, payload)


def _vendor_master_csv(path: Path) -> None:
    rows = [
        ("vendor_id", "vendor_name", "category", "payment_terms", "country", "active"),
        ("V-1001", "Northwind Office Supply", "Office supplies", "Net 30", "US", "yes"),
        ("V-1002", "Helios Cloud Services", "Cloud infrastructure", "Net 30", "US", "yes"),
        ("V-1003", "Brightline Logistics", "Freight", "Net 30", "US", "yes"),
        ("V-1004", "Quantum Print Co.", "Marketing", "Due on receipt", "US", "yes"),
        ("V-1005", "Atlas Facilities Management", "Facilities", "Net 30", "US", "yes"),
        ("V-1006", "Bluebird Catering", "Events", "Net 14", "US", "yes"),
        ("V-1007", "Cedar & Pine Legal LLP", "Legal", "Net 30", "US", "yes"),
        ("V-1008", "Sierra Components Ltd.", "Hardware", "Net 45", "CA", "yes"),
        ("V-1009", "Orion Labs LLC", "R&D partner", "n/a", "US", "no"),
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows(rows)


def generate_corpus(root: Path) -> list[Path]:
    root.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for invoice in INVOICES:
        target = root / invoice.filename
        _invoice_pdf(invoice, target)
        written.append(target)
    for contract in CONTRACTS:
        target = root / contract.filename
        _contract_docx(contract, target)
        written.append(target)
    lease = root / LEASE.filename
    _contract_pdf(LEASE, lease)
    xlsx = root / "invoices/2026-04/sierra-components-SC-90021.xlsx"
    _invoice_xlsx(xlsx)
    vendors = root / "reference/vendor-master.csv"
    _vendor_master_csv(vendors)
    broken = root / "inbox-scans/scan-2026-04-17-0042.pdf"
    broken.parent.mkdir(parents=True, exist_ok=True)
    # A truncated upload from a scanner: valid header, no body. It should land in the dead-letter queue.
    broken.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R\n")
    written.extend([lease, xlsx, vendors, broken])
    return written


def amend_invoice(root: Path) -> Path:
    """Simulates a vendor re-issuing an invoice so the next run has exactly one changed document."""
    original = INVOICES[0]
    revised = InvoiceSpec(
        filename=original.filename,
        vendor=original.vendor,
        vendor_address=original.vendor_address,
        number=original.number,
        issued=original.issued,
        due=original.due,
        items=(*original.items, LineItem("Monitor arm, dual", D("6"), D("129.00"))),
        tax_rate=original.tax_rate,
        notes="Revised: added monitor arms per PO-88213 change order 1.",
        po_number=original.po_number,
    )
    target = root / revised.filename
    _invoice_pdf(revised, target)
    return target
