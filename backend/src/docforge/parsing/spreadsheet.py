from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from docforge.domain.documents import (
    BBox,
    Block,
    BlockKind,
    DocumentFormat,
    DocumentModel,
    GeometryKind,
    Page,
    Word,
)
from docforge.errors import ParseError

CELL_WIDTH = 84.0
CELL_HEIGHT = 18.0
MARGIN = 36.0
TABLE_MIN_CELLS = 3
_FIXED_DECIMALS = re.compile(r"0\.(0+)")


def _shown_decimals(number_format: str) -> int:
    match = _FIXED_DECIMALS.search(number_format)
    return len(match.group(1)) if match else 0


def _format_number(value: float | Decimal, number_format: str) -> str:
    """Exact digits, never rounded and never in exponent form, padded to the decimals the cell
    displays: 12.5 shown as 0.00 reads 12.50, while 1.375 stays 1.375 rather than 1.38."""
    if not math.isfinite(value):
        return str(value)
    exact = Decimal(str(value)).normalize()
    exponent = exact.as_tuple().exponent
    places = max(-exponent if isinstance(exponent, int) else 0, _shown_decimals(number_format), 0)
    return format(exact.quantize(Decimal(1).scaleb(-places)), "f")


def _format_cell(value: object, number_format: str = "General") -> str:
    match value:
        case None:
            return ""
        case bool():
            return "true" if value else "false"
        case datetime():
            return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        case date():
            return value.isoformat()
        case int():
            return str(value)
        case float() | Decimal():
            return _format_number(value, number_format)
        case _:
            return " ".join(str(value).split())


def _cell_box(row: int, col: int) -> BBox:
    return BBox(
        x0=MARGIN + col * CELL_WIDTH,
        top=MARGIN + row * CELL_HEIGHT,
        x1=MARGIN + (col + 1) * CELL_WIDTH,
        bottom=MARGIN + (row + 1) * CELL_HEIGHT,
    )


def _range(first_row: int, last_row: int, first_col: int, last_col: int) -> str:
    return (
        f"{get_column_letter(first_col + 1)}{first_row + 1}:{get_column_letter(last_col + 1)}{last_row + 1}"
    )


def grid_to_page(number: int, grid: Sequence[Sequence[str]], label: str | None) -> Page:
    """Turn a cell grid into blocks: dense row runs become tables, sparse rows become text."""
    width = max((len(row) for row in grid), default=1)
    page = Page(
        number=number,
        width=max(612.0, 2 * MARGIN + width * CELL_WIDTH),
        height=max(792.0, 2 * MARGIN + len(grid) * CELL_HEIGHT),
        label=label,
    )

    def filled(row: Sequence[str]) -> list[int]:
        return [i for i, value in enumerate(row) if value]

    r = 0
    while r < len(grid):
        cols = filled(grid[r])
        if not cols:
            r += 1
            continue
        if len(cols) >= TABLE_MIN_CELLS:
            start = r
            while r < len(grid) and len(filled(grid[r])) >= 2:
                r += 1
            rows = grid[start:r]
            first_col = min(min(filled(row)) for row in rows)
            last_col = max(max(filled(row)) for row in rows)
            trimmed = [list(row[first_col : last_col + 1]) for row in rows]
            trimmed = [row + [""] * (last_col - first_col + 1 - len(row)) for row in trimmed]
            words = [
                Word(text=value, bbox=_cell_box(start + i, first_col + j))
                for i, row in enumerate(trimmed)
                for j, value in enumerate(row)
                if value
            ]
            page.blocks.append(
                Block(
                    id=f"p{number}-b{len(page.blocks)}",
                    page=number,
                    kind=BlockKind.TABLE,
                    text="\n".join(" | ".join(row) for row in trimmed),
                    bbox=_cell_box(start, first_col).union(_cell_box(r - 1, last_col)),
                    geometry=GeometryKind.GRID,
                    rows=trimmed,
                    row_boxes=[
                        _cell_box(start + i, first_col).union(_cell_box(start + i, last_col))
                        for i in range(len(trimmed))
                    ],
                    cell_range=_range(start, r - 1, first_col, last_col),
                    words=words,
                )
            )
            continue

        values = [grid[r][i] for i in cols]
        next_is_table = r + 1 < len(grid) and len(filled(grid[r + 1])) >= TABLE_MIN_CELLS
        is_heading = len(values) == 1 and (not page.blocks or next_is_table)
        if len(values) == 2 and not values[0].endswith(":"):
            text = f"{values[0]}: {values[1]}"
        else:
            text = " ".join(values)
        page.blocks.append(
            Block(
                id=f"p{number}-b{len(page.blocks)}",
                page=number,
                kind=BlockKind.HEADING if is_heading else BlockKind.PARAGRAPH,
                text=text,
                bbox=_cell_box(r, cols[0]).union(_cell_box(r, cols[-1])),
                geometry=GeometryKind.GRID,
                heading_level=(1 if not page.blocks else 2) if is_heading else None,
                cell_range=_range(r, r, cols[0], cols[-1]),
                words=[Word(text=grid[r][i], bbox=_cell_box(r, i)) for i in cols],
            )
        )
        r += 1
    return page


def _normalize(rows: Iterable[Sequence[object]]) -> list[list[str]]:
    grid = [[_format_cell(value) for value in row] for row in rows]
    return _trim(grid)


def _normalize_cells(rows: Iterable[Sequence[object]]) -> list[list[str]]:
    grid = [
        [
            _format_cell(getattr(cell, "value", None), getattr(cell, "number_format", "General"))
            for cell in row
        ]
        for row in rows
    ]
    return _trim(grid)


def _trim(grid: list[list[str]]) -> list[list[str]]:
    while grid and not any(grid[-1]):
        grid.pop()
    return grid


class XlsxParser:
    formats = frozenset({DocumentFormat.XLSX})
    version = "xlsx-2"

    def parse(self, data: bytes) -> DocumentModel:
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise ParseError(f"malformed XLSX: {exc}") from exc
        try:
            pages: list[Page] = []
            for sheet in workbook.worksheets:
                grid = _normalize_cells(sheet.iter_rows())
                if grid:
                    pages.append(grid_to_page(len(pages) + 1, grid, label=sheet.title))
        finally:
            workbook.close()
        if not pages:
            raise ParseError("workbook has no non-empty sheets")
        return DocumentModel(
            format=DocumentFormat.XLSX,
            pages=pages,
            metadata={"sheets": ",".join(p.label or "" for p in pages)},
        )


class CsvParser:
    formats = frozenset({DocumentFormat.CSV})
    version = "csv-2"

    def parse(self, data: bytes) -> DocumentModel:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ParseError(f"CSV is not valid UTF-8 (byte {exc.start}): re-export it as UTF-8") from exc
        if not text.strip():
            raise ParseError("CSV file is empty")
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        grid = _normalize(csv.reader(io.StringIO(text), dialect))
        return DocumentModel(format=DocumentFormat.CSV, pages=[grid_to_page(1, grid, label=None)])
